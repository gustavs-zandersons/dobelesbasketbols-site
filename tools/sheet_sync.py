"""Import the season sheet (Google Sheets) into the site's data files.

    python3 tools/sheet_sync.py --url <link to the published or shared sheet>
    python3 tools/sheet_sync.py --file <sheet.xlsx>
    add --check to only check the sheet (nothing is written)
    add --expect-same to fail when the sheet differs from src/sheet.json
        (CI: the template built from src/sheet.json must import back unchanged)

Reads the tabs "Iestatījumi", "Komandas", "Zāles", "Sponsori", "Kalendārs",
"MVP", "Mēneša MVP" and "Lēmumi". The site's games come from the calendar
together with basket.lv (tools/basket_sync.py, tools/real-season.js).

Writes src/sheet.json, src/sponsors.json, src/logo-sources.json and any new
logos (assets/teams/<id>.webp, assets/sponsors/<id>.webp). Logos come from the
"Logo (saite)" column: a Google Drive file link (shared "Anyone with the link")
or any https link to a PNG/JPG/WebP. An empty cell keeps the logo the site has.

Prints a report in Latvian (also to the GitHub job summary when run there).
Exit code 1 when the sheet has errors: then nothing is written, so the site
stays as it was.
"""
import argparse
import hashlib
import io
import json
import os
import re
import sys
import urllib.error
import urllib.request
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'src'
SHEET_JSON = SRC / 'sheet.json'
SPONSORS_JSON = SRC / 'sponsors.json'
SOURCES_JSON = SRC / 'logo-sources.json'

LEVELS = ('Ģenerālsponsors', 'Atbalstītājs')
GROUPS = ('A', 'B')
LINK_KEYS = ('website', 'facebook', 'instagram', 'tiktok')
ID_RE = re.compile(r'^[a-z0-9]+(-[a-z0-9]+)*$')
LINK_RE = re.compile(r'^https://\S+$')
TEAM_COUNT = 12          # the demo season needs exactly 12 teams (6 + 6)
SHORT_MAX = 12           # longer short names get cramped in tables and posters
MAX_DOWNLOAD = 10 * 1024 * 1024
TEAM_BOX, TEAM_CANVAS = 240, 256      # team logos: square canvas, logo fitted inside
# The site's own address. In the sheet, a logo link to the site itself
# (…/assets/teams/<id>.webp) means "the logo the site already has".
SITE_URL = 'https://dobelesbasketbols.lv/'
# Older sheet links may still use the github.io address (it forwards to the domain).
SITE_HOSTS = ('dobelesbasketbols.lv', 'www.dobelesbasketbols.lv', 'gustavs-zandersons.github.io')
SPONSOR_BOX = (400, 200)              # sponsor logos: same as tools/prepare_assets.py

# Iestatījumi: parameter (column A) -> key in sheet.json settings, required?, kind
SETTINGS = [
    ('Sezona', 'season', True, 'text'),
    ('DBČK nosaukums', 'cups.dbck.name', True, 'text'),
    ('JSSC nosaukums', 'cups.jssc.name', True, 'text'),
    ('DBČK MVP balvas sponsors', 'cups.dbck.mvpSponsor', False, 'text'),
    ('JSSC MVP balvas sponsors', 'cups.jssc.mvpSponsor', False, 'text'),
    ('DBČK mēneša MVP sponsors', 'cups.dbck.monthMvpSponsor', False, 'text'),
    ('JSSC mēneša MVP sponsors', 'cups.jssc.monthMvpSponsor', False, 'text'),
    ('DBČK basket.lv saite', 'cups.dbck.basket', False, 'link'),
    ('JSSC basket.lv saite', 'cups.jssc.basket', False, 'link'),
    ('Sponsoru lapas teksts', 'sponsorsText', False, 'text'),
    ('Kājenes teksts', 'footer.text', False, 'text'),
    ('Turnīra Facebook', 'footer.links.facebook', False, 'link'),
    ('Turnīra Instagram', 'footer.links.instagram', False, 'link'),
    ('Turnīra YouTube', 'footer.links.youtube', False, 'link'),
    ('Turnīra TikTok', 'footer.links.tiktok', False, 'link'),
]

# Columns per tab: key, header(s) in the sheet, required column?
LOGO_HEADERS = ('Logo (saite)', 'Logo saite', 'Logo', 'Logo fails')
TEAM_COLS = [
    ('id', ('ID',), True), ('name', ('Pilnais nosaukums',), True), ('short', ('Īsais nosaukums',), True),
    ('basket', ('basket.lv nosaukums',), False), ('group', ('Grupa',), True), ('logo', LOGO_HEADERS, False),
    ('city', ('Pilsēta',), False), ('description', ('Apraksts',), False), ('website', ('Mājaslapa',), False),
    ('facebook', ('Facebook',), False), ('instagram', ('Instagram',), False), ('tiktok', ('TikTok',), False),
]
HALL_COLS = [('id', ('ID',), True), ('name', ('Nosaukums',), True), ('address', ('Adrese',), False)]
SPONSOR_COLS = [
    ('id', ('ID',), True), ('name', ('Nosaukums',), True), ('level', ('Līmenis',), True), ('order', ('Secība',), True),
    ('description', ('Apraksts',), False), ('website', ('Mājaslapa',), False), ('facebook', ('Facebook',), False),
    ('instagram', ('Instagram',), False), ('tiktok', ('TikTok',), False), ('logo', LOGO_HEADERS, False),
]
LINK_NAMES = {'website': 'Mājaslapa', 'facebook': 'Facebook', 'instagram': 'Instagram', 'tiktok': 'TikTok',
              'logo': 'Logo (saite)'}

# Kalendārs, MVP, Mēneša MVP, Lēmumi
CAL_COLS = [
    ('id', ('Spēles ID',), True), ('cup', ('Kauss',), True), ('stage', ('Posms',), True), ('round', ('Kārta',), False),
    ('tie', ('Pāris',), False), ('n', ('Spēle Nr.',), False), ('date', ('Datums',), True), ('time', ('Laiks',), True),
    ('hall', ('Zāle',), True), ('home', ('Mājinieki',), True), ('away', ('Viesi',), True), ('status', ('Statuss',), False),
    ('basket', ('basket.lv saite',), False), ('note', ('Piezīme',), False),
]
MVP_COLS = [('cup', ('Kauss',), True), ('date', ('Kārtas datums',), True), ('player', ('Spēlētājs',), True),
            ('team', ('Komanda',), True), ('game', ('Spēles ID',), False), ('note', ('Piezīme',), False)]
MONTH_COLS = [('cup', ('Kauss',), True), ('month', ('Mēnesis',), True), ('player', ('Spēlētājs',), True),
              ('team', ('Komanda',), True), ('note', ('Piezīme',), False)]
DEC_COLS = [('game', ('Spēles ID',), True), ('kind', ('Veids',), True), ('value', ('Vērtība',), False),
            ('date', ('Lēmuma datums',), False), ('by', ('Kas pieņēma',), False), ('note', ('Piezīme',), False)]
CUP_CODES = {'DBČK': 'dbck', 'JSSC': 'jssc'}
CUP_LABELS = {v: k for k, v in CUP_CODES.items()}
KG = 'Playoff kvalifikācijas grupa'
# Posms per cup -> the site's game stage
STAGES = {
    'dbck': {'Grupu turnīrs': 'group', KG: 'k', 'Cerību kauss': 'hope', 'Pusfināls': 'playoff',
             'Spēle par 3. vietu': 'playoff', 'Fināls': 'playoff'},
    'jssc': {'Astotdaļfināls': 'cup', 'Ceturtdaļfināls': 'cup', 'Pusfināls': 'cup', 'Fināls': 'cup'},
}
# Pairs (Pāris) per cup -> the Posms they belong to
TIES = {
    'dbck': {'H1': 'Cerību kauss', 'H2': 'Cerību kauss', 'HF': 'Cerību kauss', 'SF1': 'Pusfināls', 'SF2': 'Pusfināls',
             'BRONZE': 'Spēle par 3. vietu', 'GOLD': 'Fināls'},
    'jssc': dict([(f'R16_{i}', 'Astotdaļfināls') for i in range(1, 5)] + [(f'QF{i}', 'Ceturtdaļfināls') for i in range(1, 5)] +
                 [('SF1', 'Pusfināls'), ('SF2', 'Pusfināls'), ('F', 'Fināls')]),
}
STATUSES = ('Plānota', 'Ja vajag', 'Pārcelta', 'Atcelta')
SLOT_RE = re.compile(r'^(?:[AB][1-6]|K[1-8]|[WL]:[A-Z0-9_]+)$')
GAME_ID_RE = re.compile(r'^[A-Za-z0-9_]+(-[A-Za-z0-9_]+)*$')
DECISIONS = ('Neieradās mājinieki', 'Neieradās viesi', 'Rezultāts anulēts', 'Rezultāta labojums', 'Izloze')
MONTHS = ('Janvāris', 'Februāris', 'Marts', 'Aprīlis', 'Maijs', 'Jūnijs', 'Jūlijs', 'Augusts', 'Septembris',
          'Oktobris', 'Novembris', 'Decembris')


class SyncError(Exception):
    """The sheet could not be read at all (bad link, not published, …)."""


class Report:
    def __init__(self):
        self.errors, self.warnings, self.changes = [], [], []

    def err(self, tab, row, msg):
        self.errors.append((tab, row, msg))

    def warn(self, tab, row, msg):
        self.warnings.append((tab, row, msg))

    def markdown(self):
        out = []
        if self.errors:
            out.append(f'## Tabulā ir kļūdas ({len(self.errors)}) — mājaslapa netika atjaunota\n')
            out.append('Izlabo tās tabulā un palaid atjaunošanu vēlreiz.\n')
            out += table(self.errors, 'Kas jālabo')
        else:
            out.append('## Tabula nolasīta bez kļūdām\n')
            if self.changes:
                out.append('**Kas mainījās:**\n')
                out += [f'- {c}' for c in self.changes]
            else:
                out.append('Izmaiņu nav — mājaslapā jau ir jaunākie dati.')
        if self.warnings:
            out.append(f'\n### Brīdinājumi ({len(self.warnings)}) — netraucē atjaunošanai\n')
            out += table(self.warnings, 'Kas jāzina')
        return '\n'.join(out) + '\n'


TAB_ORDER = ('Iestatījumi', 'Komandas', 'Zāles', 'Sponsori', 'Kalendārs', 'MVP', 'Mēneša MVP', 'Lēmumi')


def table(items, what):
    rows = [f'| Lapa | Rinda | {what} |', '|---|---|---|']
    items = sorted(items, key=lambda x: (TAB_ORDER.index(x[0]) if x[0] in TAB_ORDER else 9, x[1] or 0))
    for tab, row, msg in items:
        rows.append(f'| {tab} | {row or "—"} | {msg.replace("|", "/")} |')
    return rows


# ---------------------------------------------------------------- download

def sheet_export_url(url):
    """Any link to the sheet -> a link that downloads the whole workbook (.xlsx)."""
    m = re.search(r'/spreadsheets/d/e/([A-Za-z0-9_-]+)', url)
    if m:   # File -> Share -> Publish to web
        return f'https://docs.google.com/spreadsheets/d/e/{m.group(1)}/pub?output=xlsx'
    m = re.search(r'/spreadsheets/d/([A-Za-z0-9_-]{20,})', url)
    if m:   # shared with "Anyone with the link"
        return f'https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx'
    raise SyncError('Tabulas saite nav atpazīta. Vajag Google Sheets saiti '
                    '(Fails → Kopīgot → Publicēt tīmeklī → „Viss dokuments”).')


def drive_download_url(url):
    """A Google Drive file link -> a direct download link. Other links unchanged."""
    if 'drive.google.com' not in url and 'docs.google.com' not in url:
        return url
    if '/folders/' in url:
        raise ValueError('šī ir mapes saite — vajag konkrētā logo faila saiti')
    m = re.search(r'/file/d/([A-Za-z0-9_-]+)', url) or re.search(r'[?&]id=([A-Za-z0-9_-]+)', url)
    if not m:
        raise ValueError('Google Drive saite nav atpazīta — kopē saiti no faila: Kopīgot → Kopēt saiti')
    return f'https://drive.google.com/uc?export=download&id={m.group(1)}'


def fetch(url, what):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (dobelesbasketbols sheet sync)'})
    try:
        with urllib.request.urlopen(req, timeout=40) as r:
            body = r.read(MAX_DOWNLOAD + 1)
            ctype = r.headers.get('Content-Type', '')
    except urllib.error.HTTPError as e:
        raise ValueError(f'{what}: serveris atbildēja ar kļūdu {e.code}')
    except (urllib.error.URLError, TimeoutError, OSError) as e:
        raise ValueError(f'{what}: neizdevās lejupielādēt ({getattr(e, "reason", e)})')
    if len(body) > MAX_DOWNLOAD:
        raise ValueError(f'{what}: fails ir lielāks par 10 MB')
    return body, ctype


def allowed_link(url):
    if LINK_RE.match(url):
        return True
    # Local test server only (tools/test: SHEET_SYNC_TEST=1).
    return os.environ.get('SHEET_SYNC_TEST') == '1' and url.startswith('http://127.0.0.1')


# ---------------------------------------------------------------- reading

def text(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return v.strftime('%d.%m.%Y')
    return re.sub(r'[ \t ]+', ' ', str(v)).strip()


def norm(h):
    return text(h).casefold()


def read_tab(wb, name, cols, rep, raw=()):
    """Rows of one tab as (row number, {key: text}). Columns are found by header.
    Keys in `raw` keep the cell's own value (dates and times)."""
    if name not in wb.sheetnames:
        rep.err(name, None, f'Tabulā nav lapas „{name}”. Lapas nosaukumu mainīt nedrīkst.')
        return []
    ws = wb[name]
    head = [norm(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    where, missing = {}, False
    for key, titles, required in cols:
        idx = next((head.index(t.casefold()) for t in titles if t.casefold() in head), None)
        if idx is None and required:
            rep.err(name, 1, f'Nav kolonnas „{titles[0]}”. Kolonnu nosaukumus mainīt nedrīkst.')
            missing = True
        if idx is not None:
            where[key] = idx
    if missing:          # the rows can't be read reliably until the header is fixed
        return []
    rows = []
    for r, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = {k: (cells[i] if k in raw else text(cells[i])) if i < len(cells) else '' for k, i in where.items()}
        for k in raw:
            if isinstance(vals.get(k), str):
                vals[k] = text(vals[k])
            elif vals.get(k) is None:
                vals[k] = ''
        if any(v not in ('', None) for v in vals.values()):
            rows.append((r, vals))
    return rows


def check_links(tab, r, vals, keys, rep):
    for k in keys:
        v = vals.get(k, '')
        if v and not allowed_link(v):
            rep.err(tab, r, f'{LINK_NAMES[k]}: saitei jāsākas ar https:// (tagad: „{v[:60]}”).')


def check_id(tab, r, v, seen, rep, what):
    if not v:
        rep.err(tab, r, f'Rindā ir dati, bet nav ID. Ieraksti {what} ID vai izdzēs rindu.')
        return False
    if not ID_RE.match(v):
        rep.err(tab, r, f'ID „{v}”: drīkst lietot tikai mazos latīņu burtus bez garumzīmēm, ciparus un domuzīmi (piem., „buki”, „ozols-pub”).')
        return False
    if v in seen:
        rep.err(tab, r, f'ID „{v}” jau ir {seen[v]}. rindā — katram ID jābūt atšķirīgam.')
        return False
    seen[v] = r
    return True


def read_settings(wb, rep):
    tab = 'Iestatījumi'
    rows = read_tab(wb, tab, [('param', ('Parametrs',), True), ('value', ('Vērtība',), True)], rep)
    found = {}
    for r, v in rows:
        if v.get('param'):
            found[v['param'].casefold()] = (r, v.get('value', ''), v['param'])
    known = {p.casefold() for p, *_ in SETTINGS}
    for p, (r, _, original) in found.items():
        if p not in known:
            rep.warn(tab, r, f'Nezināms parametrs „{original}” — mājaslapa to neizmanto.')
    out = {}
    for param, key, required, kind in SETTINGS:
        r, val, _ = found.get(param.casefold(), (None, '', ''))
        if r is None and required:
            rep.err(tab, None, f'Nav parametra „{param}”. Parametru nosaukumus mainīt nedrīkst.')
        if required and r is not None and not val:
            rep.err(tab, r, f'„{param}” jāaizpilda.')
        if kind == 'link' and val and not allowed_link(val):
            rep.err(tab, r, f'„{param}”: saitei jāsākas ar https:// (tagad: „{val[:60]}”).')
        node = out
        parts = key.split('.')
        for p in parts[:-1]:
            node = node.setdefault(p, {})
        node[parts[-1]] = val
    if out.get('season') and not re.match(r'^\d{4}/\d{2}$', out['season']):
        rep.warn(tab, found.get('sezona', (None,))[0], f'Sezonu parasti raksta kā „2026/27” (tagad: „{out["season"]}”).')
    return out


def read_teams(wb, rep):
    tab = 'Komandas'
    seen, teams = {}, []
    for r, v in read_tab(wb, tab, TEAM_COLS, rep):
        if not check_id(tab, r, v.get('id', ''), seen, rep, 'komandas'):
            continue
        if not v.get('name'):
            rep.err(tab, r, 'Nav pilnā nosaukuma.')
        if not v.get('short'):
            rep.err(tab, r, 'Nav īsā nosaukuma (to rāda tabulās un plakātos).')
        elif len(v['short']) > SHORT_MAX:
            rep.warn(tab, r, f'Īsais nosaukums „{v["short"]}” ir garāks par {SHORT_MAX} zīmēm — tabulās un plakātos var neietilpt.')
        g = v.get('group', '').upper()
        if g and g not in GROUPS:
            rep.err(tab, r, f'Grupa var būt tikai A vai B (tagad: „{v["group"]}”).')
        check_links(tab, r, v, LINK_KEYS + ('logo',), rep)
        t = {'id': v['id'], 'name': v.get('name', ''), 'short': v.get('short', ''), 'group': g if g in GROUPS else ''}
        aliases = [a.strip() for a in v.get('basket', '').split(';') if a.strip()]
        if aliases:
            t['basket'] = aliases
        for k in ('city', 'description') + LINK_KEYS:
            if v.get(k):
                t[k] = v[k]
        teams.append((r, t, v.get('logo', '')))
    if len(teams) != TEAM_COUNT and not any(e[0] == tab for e in rep.errors):
        rep.err(tab, None, f'Ir {len(teams)} komandas, bet sezonai (un demo kalendāram) vajag tieši {TEAM_COUNT}.')
    groups = [t['group'] for _, t, _ in teams]
    if any(groups):
        if groups.count('A') != 6 or groups.count('B') != 6:
            rep.warn(tab, None, f'Grupas aizpildītas daļēji (A: {groups.count("A")}, B: {groups.count("B")}). '
                                'Kamēr nav 6 + 6, demo kalendārs izmanto izlozes simulāciju.')
    return teams


def read_halls(wb, rep):
    tab = 'Zāles'
    seen, halls = {}, []
    for r, v in read_tab(wb, tab, HALL_COLS, rep):
        if not check_id(tab, r, v.get('id', ''), seen, rep, 'zāles'):
            continue
        if not v.get('name'):
            rep.err(tab, r, 'Nav zāles nosaukuma.')
        h = {'id': v['id'], 'name': v.get('name', '')}
        if v.get('address'):
            h['address'] = v['address']
        halls.append(h)
    if not halls and not any(e[0] == tab for e in rep.errors):
        rep.err(tab, None, 'Nav nevienas zāles.')
    elif len(halls) == 1:
        rep.warn(tab, None, 'Ir tikai viena zāle — demo kalendārs visas spēles liks tajā.')
    return halls


def read_sponsors(wb, rep):
    tab = 'Sponsori'
    seen, out = {}, []
    for r, v in read_tab(wb, tab, SPONSOR_COLS, rep):
        if not check_id(tab, r, v.get('id', ''), seen, rep, 'sponsora'):
            continue
        if not v.get('name'):
            rep.err(tab, r, 'Nav nosaukuma.')
        if v.get('level') not in LEVELS:
            rep.err(tab, r, f'Līmenim jābūt „{LEVELS[0]}” vai „{LEVELS[1]}” (tagad: „{v.get("level", "")}”).')
        order = v.get('order', '')
        if order and not re.match(r'^\d+$', order):
            rep.err(tab, r, f'Secībai jābūt skaitlim (tagad: „{order}”).')
        check_links(tab, r, v, LINK_KEYS + ('logo',), rep)
        s = {'id': v['id'], 'name': v.get('name', ''), 'level': v.get('level', ''),
             'order': int(order) if order.isdigit() else 999, 'description': v.get('description', '')}
        for k in LINK_KEYS:
            s[k] = v.get(k, '')
        out.append((r, s, v.get('logo', '')))
    if not out and not any(e[0] == tab for e in rep.errors):
        rep.err(tab, None, 'Nav neviena sponsora.')
    return out


# ---------------------------------------------------------------- calendar, MVPs, decisions

def parse_date(v):
    """A cell -> 'YYYY-MM-DD', '' for empty, None when it can't be read."""
    if v in ('', None):
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{4})\.?$', s) or re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})$', s)
    if not m:
        return None
    a, b, c = (int(x) for x in m.groups())
    y, mo, d = (a, b, c) if a > 31 else (c, b, a)
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


def parse_time(v):
    """A cell -> 'HH:MM', '' for empty, None when it can't be read."""
    from datetime import time as dtime
    if v in ('', None):
        return ''
    if isinstance(v, datetime):
        v = v.time()
    if isinstance(v, dtime):
        return f'{v.hour:02d}:{v.minute:02d}'
    if isinstance(v, (int, float)) and 0 <= v < 1:        # a time stored as a fraction of a day
        mins = round(v * 24 * 60)
        return f'{mins // 60:02d}:{mins % 60:02d}'
    m = re.match(r'^(\d{1,2})[:.](\d{2})(?::\d{2})?$', str(v).strip())
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        return None
    return f'{int(m.group(1)):02d}:{m.group(2)}'


def read_calendar(wb, rep, team_ids, hall_ids, season_years):
    tab = 'Kalendārs'
    seen, out = {}, []
    lo, hi = (f'{season_years[0]}-08-01', f'{season_years[1]}-07-31') if season_years else ('0000', '9999')
    for r, v in read_tab(wb, tab, CAL_COLS, rep, raw=('date', 'time')):
        gid = v.get('id', '')
        if not gid:
            rep.err(tab, r, 'Rindā ir dati, bet nav spēles ID. Ieraksti ID vai izdzēs rindu.')
            continue
        if not GAME_ID_RE.match(gid):
            rep.err(tab, r, f'Spēles ID „{gid}”: drīkst lietot tikai latīņu burtus, ciparus, domuzīmi un apakšsvītru (piem., „DBCK-GR-01”).')
            continue
        if gid.casefold() in seen:
            rep.err(tab, r, f'Spēles ID „{gid}” jau ir {seen[gid.casefold()]}. rindā — katrai spēlei savs ID.')
            continue
        seen[gid.casefold()] = r
        cup = CUP_CODES.get(v.get('cup', ''))
        if not cup:
            rep.err(tab, r, f'{gid}: kausam jābūt „DBČK” vai „JSSC” (tagad: „{v.get("cup", "")}”).')
            continue
        stage = v.get('stage', '')
        if stage not in STAGES[cup]:
            rep.err(tab, r, f'{gid}: posms „{stage}” nav {v["cup"]} posmu sarakstā (izvēlies no saraksta).')
            continue
        rnd = None
        if v.get('round'):
            m = re.match(r'^(\d+)\.?\s*(kārta)?$', v['round'])
            if m:
                rnd = int(m.group(1))
            else:
                rep.err(tab, r, f'{gid}: kārta jāraksta kā „1. kārta” (tagad: „{v["round"]}”).')
        tie = v.get('tie', '')
        if STAGES[cup][stage] in ('group', 'k'):
            if tie:
                rep.warn(tab, r, f'{gid}: grupu spēlēm pāri nenorāda — „{tie}” netiek izmantots.')
                tie = ''
            if rnd is None and parse_date(v.get('date')):
                rep.warn(tab, r, f'{gid}: nav norādīta kārta.')
        else:
            if tie not in TIES[cup]:
                rep.err(tab, r, f'{gid}: izslēgšanas spēlei jānorāda pāris no saraksta (tagad: „{tie}”).')
                continue
            if TIES[cup][tie] != stage:
                rep.err(tab, r, f'{gid}: pāris {tie} pieder posmam „{TIES[cup][tie]}”, nevis „{stage}”.')
        n = None
        if v.get('n'):
            if re.match(r'^[1-3]$', v['n']):
                n = int(v['n'])
            else:
                rep.err(tab, r, f'{gid}: spēles numurs var būt 1, 2 vai 3 (tagad: „{v["n"]}”).')
        d = parse_date(v.get('date'))
        if d is None:
            rep.err(tab, r, f'{gid}: datumu neizdevās nolasīt („{text(v.get("date"))}”). Raksti DD.MM.GGGG.')
            d = ''
        elif d and not lo <= d <= hi:
            rep.err(tab, r, f'{gid}: datums {d} nav šajā sezonā.')
        t = parse_time(v.get('time'))
        if t is None:
            rep.err(tab, r, f'{gid}: laiku neizdevās nolasīt („{text(v.get("time"))}”). Raksti, piem., 11:00.')
            t = ''
        hall = v.get('hall', '')
        if hall and hall not in hall_ids:
            rep.err(tab, r, f'{gid}: zāle „{hall}” nav lapā „Zāles” (norādi zāles ID).')
        elif d and not hall:
            rep.warn(tab, r, f'{gid}: nav norādīta zāle.')
        sides = []
        for key, label in (('home', 'Mājinieki'), ('away', 'Viesi')):
            x = v.get(key, '')
            if not x:
                rep.err(tab, r, f'{gid}: nav norādīti {label.lower()} (komandas ID vai vietas kods, piem., A1).')
            elif x not in team_ids and not SLOT_RE.match(x):
                rep.err(tab, r, f'{gid}: {label.lower()} „{x}” — nav ne komandas ID, ne vietas kods (A1…B6, K1…K8, W:SF1 …).')
            elif STAGES[cup][stage] == 'group' and x not in team_ids:
                rep.err(tab, r, f'{gid}: grupu turnīrā jānorāda komandas ID (tagad: „{x}”).')
            sides.append(x)
        if sides[0] and sides[0] == sides[1]:
            rep.err(tab, r, f'{gid}: mājinieki un viesi ir viena komanda.')
        status = v.get('status', '') or 'Plānota'
        if status not in STATUSES:
            rep.err(tab, r, f'{gid}: statuss „{status}” — izvēlies no saraksta ({", ".join(STATUSES)}).')
        link = v.get('basket', '')
        if link and not allowed_link(link):
            rep.err(tab, r, f'{gid}: basket.lv saitei jāsākas ar https://.')
        out.append({'id': gid, 'cup': cup, 'stage': stage, 'round': rnd, 'tie': tie, 'n': n, 'date': d, 'time': t,
                    'hall': hall, 'home': sides[0], 'away': sides[1], 'status': status, 'basket': link,
                    'note': v.get('note', '')})
    return out


def read_mvps(wb, rep, team_ids, game_ids):
    tab = 'MVP'
    out = []
    for r, v in read_tab(wb, tab, MVP_COLS, rep, raw=('date',)):
        cup = CUP_CODES.get(v.get('cup', ''))
        d = parse_date(v.get('date'))
        ok = True
        if not cup:
            rep.err(tab, r, f'Kausam jābūt „DBČK” vai „JSSC” (tagad: „{v.get("cup", "")}”).'); ok = False
        if not d:
            rep.err(tab, r, 'Kārtas datumu neizdevās nolasīt. Raksti DD.MM.GGGG.'); ok = False
        if not v.get('player'):
            rep.err(tab, r, 'Nav spēlētāja vārda.'); ok = False
        if v.get('team') not in team_ids:
            rep.err(tab, r, f'Komanda „{v.get("team", "")}” — norādi komandas ID no saraksta.'); ok = False
        game = v.get('game', '')
        if game and game.casefold() not in game_ids:
            rep.warn(tab, r, f'Spēle „{game}” nav kalendārā — spēli atradīs pēc datuma un komandas.')
            game = ''
        if ok:
            out.append({'cup': cup, 'date': d, 'player': v['player'], 'team': v['team'], 'game': game, 'note': v.get('note', '')})
    return out


def read_month_mvps(wb, rep, team_ids):
    tab = 'Mēneša MVP'
    out, seen = [], {}
    names = {m.casefold(): i + 1 for i, m in enumerate(MONTHS)}
    for r, v in read_tab(wb, tab, MONTH_COLS, rep):
        cup = CUP_CODES.get(v.get('cup', ''))
        m = re.match(r'^(\S+)\s+(\d{4})$', v.get('month', ''))
        month = f'{m.group(2)}-{names[m.group(1).casefold()]:02d}' if m and m.group(1).casefold() in names else None
        ok = True
        if not cup:
            rep.err(tab, r, f'Kausam jābūt „DBČK” vai „JSSC” (tagad: „{v.get("cup", "")}”).'); ok = False
        if not month:
            rep.err(tab, r, f'Mēnesi neizdevās nolasīt („{v.get("month", "")}”) — izvēlies no saraksta.'); ok = False
        if not v.get('player'):
            rep.err(tab, r, 'Nav spēlētāja vārda.'); ok = False
        if v.get('team') not in team_ids:
            rep.err(tab, r, f'Komanda „{v.get("team", "")}” — norādi komandas ID no saraksta.'); ok = False
        if ok and (cup, month) in seen:
            rep.err(tab, r, f'{v["cup"]} {v["month"]}: MVP jau ir {seen[(cup, month)]}. rindā.'); ok = False
        if ok:
            seen[(cup, month)] = r
            out.append({'cup': cup, 'month': month, 'player': v['player'], 'team': v['team'], 'note': v.get('note', '')})
    return out


def read_decisions(wb, rep, game_ids, team_ids):
    tab = 'Lēmumi'
    out = []
    for r, v in read_tab(wb, tab, DEC_COLS, rep, raw=('date',)):
        game, kind, value = v.get('game', ''), v.get('kind', ''), v.get('value', '')
        ok = True
        if game.casefold() not in game_ids:
            rep.err(tab, r, f'Spēle „{game}” nav kalendārā (norādi spēles ID).'); ok = False
        if kind not in DECISIONS:
            rep.err(tab, r, f'Lēmuma veids „{kind}” — izvēlies no saraksta.'); ok = False
        if kind == 'Rezultāta labojums' and not re.match(r'^\d{1,3}\s*:\s*\d{1,3}$', value):
            rep.err(tab, r, f'Labojumam vērtībā jāraksta rezultāts, piem., 78:71 (tagad: „{value}”).'); ok = False
        if kind == 'Izloze':
            ids = [x.strip() for x in value.split(',') if x.strip()]
            bad = [x for x in ids if x not in team_ids]
            if len(ids) < 2 or bad:
                rep.err(tab, r, f'Izlozei vērtībā jāraksta komandu ID izlozes secībā, piem., „buki, doks” (tagad: „{value}”).'); ok = False
        d = parse_date(v.get('date'))
        if d is None:
            rep.warn(tab, r, 'Lēmuma datumu neizdevās nolasīt.'); d = ''
        if ok:
            out.append({'game': game, 'kind': kind, 'value': value, 'date': d, 'by': v.get('by', ''), 'note': v.get('note', '')})
    return out


# ---------------------------------------------------------------- logos

def trim(im, pad=0.02):
    """Crop to the visible logo (alpha > 8), keeping a small even margin."""
    alpha = im.getchannel('A').point(lambda v: 255 if v > 8 else 0)
    box = alpha.getbbox()
    if not box:
        return im
    x0, y0, x1, y1 = box
    p = round(max(x1 - x0, y1 - y0) * pad)
    return im.crop((max(0, x0 - p), max(0, y0 - p), min(im.width, x1 + p), min(im.height, y1 + p)))


def make_logo(raw, kind):
    """Downloaded bytes -> (web copy as WebP bytes, width, height, has transparency)."""
    if raw.lstrip()[:5].lower() in (b'<?xml', b'<svg ') or b'<svg' in raw[:400].lower():
        raise ValueError('SVG fails nav atbalstīts — ieliec PNG (ar caurspīdīgu fonu) vai JPG')
    try:
        im = Image.open(io.BytesIO(raw))
        im.load()
    except Exception:
        if raw[:200].lstrip().lower().startswith((b'<!doctype html', b'<html')):
            raise ValueError('saite atver mājaslapu, nevis attēlu. Google Drive failam vai mapei iestati piekļuvi '
                             '„Ikviens, kam ir saite” (Kopīgot → Vispārējā piekļuve)')
        raise ValueError('fails nav attēls (der PNG, JPG vai WebP)')
    im = ImageOps.exif_transpose(im).convert('RGBA')
    transparent = im.getchannel('A').getextrema()[0] < 250
    im = trim(im)
    if kind == 'teams':
        im.thumbnail((TEAM_BOX, TEAM_BOX), Image.LANCZOS)
        canvas = Image.new('RGBA', (TEAM_CANVAS, TEAM_CANVAS), (0, 0, 0, 0))
        canvas.paste(im, ((TEAM_CANVAS - im.width) // 2, (TEAM_CANVAS - im.height) // 2), im)
        im = canvas
    else:
        im.thumbnail(SPONSOR_BOX, Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, 'WEBP', quality=86, method=6)
    return buf.getvalue(), im.width, im.height, transparent


def own_logo(link, kind, item_id):
    """A link to a logo on the site itself: 'same' (this row's own logo),
    'other' (another team's or sponsor's), None (not a site link)."""
    m = re.match(r'^https?://([^/]+)/(?:[^?#]*/)?assets/(teams|sponsors)/([a-z0-9-]+)\.webp(?:[?#].*)?$', link)
    if not m or m.group(1) not in SITE_HOSTS:
        return None
    return 'same' if (m.group(2), m.group(3)) == (kind, item_id) else 'other'


def site_logo_link(kind, item_id):
    return f'{SITE_URL}assets/{kind}/{item_id}.webp'


def get_logos(kind, rows, sources, rep, tab):
    """Download the logos linked in the sheet. Returns {id: (webp bytes, w, h)} for
    new or changed logos, and the updated sources map."""
    new, out_sources = {}, dict(sources)
    for r, item, link in rows:
        if not link:
            continue
        key = f'{kind}/{item["id"]}'
        dest = ROOT / 'assets' / kind / f'{item["id"]}.webp'
        own = own_logo(link, kind, item['id'])
        if own == 'other':
            rep.err(tab, r, f'{item["id"]}: logo saite norāda uz cita logo mājaslapā. Atstāj šīs rindas saiti vai ielīmē jaunu Google Drive saiti.')
            continue
        if own == 'same':
            if not dest.exists():
                (rep.warn if kind == 'teams' else rep.err)(tab, r, f'{item["id"]}: mājaslapā šī logo nav. Ielīmē Google Drive saiti.')
            continue                          # the logo the site already has
        try:
            body, _ = fetch(drive_download_url(link), 'Logo')
            sha = hashlib.sha256(body).hexdigest()
            prev = sources.get(key, {})
            if prev.get('sha256') == sha and dest.exists():
                continue                      # same file as last time
            webp, w, h, transparent = make_logo(body, kind)
            new[item['id']] = (webp, w, h)
            out_sources[key] = {'url': link, 'sha256': sha}
            if not transparent:
                rep.warn(tab, r, f'{item["id"]}: logo nav ar caurspīdīgu fonu — uz tumšā fona būs redzams balts taisnstūris. Labāk PNG ar caurspīdīgu fonu.')
        except ValueError as e:
            if dest.exists():
                rep.warn(tab, r, f'{item["id"]}: logo neizdevās ielādēt ({e}). Paliek iepriekšējais logo.')
            elif kind == 'teams':
                rep.warn(tab, r, f'{item["id"]}: logo neizdevās ielādēt ({e}). Līdz tam rādīs iniciāļus.')
            else:
                rep.err(tab, r, f'{item["id"]}: logo neizdevās ielādēt ({e}).')
    return new, out_sources


# ---------------------------------------------------------------- compare + write

def describe_changes(old, new, rep):
    if not old:
        rep.changes.append('Pirmā ielāde no tabulas.')
        return
    if old.get('settings') != new['settings']:
        diff = [p for p, key, *_ in SETTINGS if dig(old.get('settings', {}), key) != dig(new['settings'], key)]
        rep.changes.append('Iestatījumi: ' + ', '.join(diff) if diff else 'Iestatījumi mainīti.')
    for key, label in (('teams', 'Komandas'), ('venues', 'Zāles'), ('sponsors', 'Sponsori')):
        a = {x['id']: x for x in old.get(key, [])}
        b = {x['id']: x for x in new[key]}
        added = [i for i in b if i not in a]
        removed = [i for i in a if i not in b]
        changed = [i for i in b if i in a and a[i] != b[i]]
        parts = []
        if added:
            parts.append('pievienotas ' + ', '.join(added))
        if removed:
            parts.append('izņemtas ' + ', '.join(removed))
        if changed:
            parts.append('mainītas ' + ', '.join(changed))
        if parts:
            rep.changes.append(f'{label}: ' + '; '.join(parts) + '.')
    cal_old = {g['id']: g for g in old.get('calendar', [])}
    cal_new = {g['id']: g for g in new['calendar']}
    if cal_old != cal_new:
        changed = [i for i in cal_new if cal_old.get(i) != cal_new[i]]
        removed = [i for i in cal_old if i not in cal_new]
        dated = sum(1 for g in new['calendar'] if g['date'])
        msg = f'Kalendārs: {len(new["calendar"])} spēles ({dated} ar datumu)'
        if changed:
            msg += '; mainītas ' + (', '.join(changed) if len(changed) <= 8 else f'{len(changed)} rindas')
        if removed:
            msg += '; izņemtas ' + ', '.join(removed)
        rep.changes.append(msg + '.')
    for key, label in (('mvps', 'MVP'), ('monthMvps', 'Mēneša MVP'), ('decisions', 'Lēmumi')):
        if old.get(key, []) != new[key]:
            rep.changes.append(f'{label}: {len(new[key])} ieraksti.')


def dig(d, key):
    for p in key.split('.'):
        d = d.get(p, {}) if isinstance(d, dict) else {}
    return d if d != {} else ''


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, indent=1) + '\n'


def sync(wb, write=True):
    rep = Report()
    settings = read_settings(wb, rep)
    teams = read_teams(wb, rep)
    halls = read_halls(wb, rep)
    sponsors = read_sponsors(wb, rep)

    names = {s['name'].casefold() for _, s, _ in sponsors}
    for cup in ('dbck', 'jssc'):
        for key, what in (('mvpSponsor', 'MVP balvas sponsors'), ('monthMvpSponsor', 'Mēneša MVP sponsors')):
            sp = dig(settings, f'cups.{cup}.{key}')
            if sp and sp.casefold() not in names:
                rep.warn('Iestatījumi', None, f'{what} „{sp}” nav lapā „Sponsori” — MVP blokā rādīs tikai nosaukumu.')

    team_ids = {t['id'] for _, t, _ in teams}
    hall_ids = {h['id'] for h in halls}
    m = re.match(r'^(\d{4})\s*/\s*(\d{2}|\d{4})$', settings.get('season', ''))
    years = (int(m.group(1)), int(m.group(2)) if len(m.group(2)) == 4 else int(m.group(1)) // 100 * 100 + int(m.group(2))) if m else None
    calendar = read_calendar(wb, rep, team_ids, hall_ids, years)
    game_ids = {g['id'].casefold() for g in calendar}
    mvps = read_mvps(wb, rep, team_ids, game_ids)
    month_mvps = read_month_mvps(wb, rep, team_ids)
    decisions = read_decisions(wb, rep, game_ids, team_ids)

    sources = json.loads(SOURCES_JSON.read_text()) if SOURCES_JSON.exists() else {}
    team_logos, sources = get_logos('teams', teams, sources, rep, 'Komandas')
    sp_logos, sources = get_logos('sponsors', sponsors, sources, rep, 'Sponsori')

    for r, t, link in teams:
        if (not link or own_logo(link, 'teams', t['id'])) and t['id'] not in team_logos and not (ROOT / 'assets' / 'teams' / f'{t["id"]}.webp').exists():
            rep.warn('Komandas', r, f'{t["id"]}: nav logo — rādīs iniciāļus. Ieliec saiti kolonnā „Logo (saite)”.')
    for r, s, link in sponsors:
        if (not link or own_logo(link, 'sponsors', s['id'])) and s['id'] not in sp_logos and not (ROOT / 'assets' / 'sponsors' / f'{s["id"]}.webp').exists():
            rep.err('Sponsori', r, f'{s["id"]}: nav logo. Ieliec saiti kolonnā „Logo (saite)”.')

    data = {'settings': settings,
            'teams': [t for _, t, _ in teams],
            'venues': halls,
            'sponsors': sorted((s for _, s, _ in sponsors),
                               key=lambda s: (LEVELS.index(s['level']) if s['level'] in LEVELS else 9, s['order'])),
            'calendar': calendar,
            'mvps': mvps,
            'monthMvps': month_mvps,
            'decisions': decisions}
    if rep.errors:
        return rep, None
    old = json.loads(SHEET_JSON.read_text()) if SHEET_JSON.exists() else None
    describe_changes(old, data, rep)
    if team_logos:
        rep.changes.append('Jauni komandu logo: ' + ', '.join(team_logos) + '.')
    if sp_logos:
        rep.changes.append('Jauni sponsoru logo: ' + ', '.join(sp_logos) + '.')

    if write:
        for kind, logos in (('teams', team_logos), ('sponsors', sp_logos)):
            for i, (webp, *_rest) in logos.items():
                d = ROOT / 'assets' / kind / f'{i}.webp'
                d.parent.mkdir(parents=True, exist_ok=True)
                d.write_bytes(webp)
        SHEET_JSON.write_text(dump(data))
        SOURCES_JSON.write_text(dump(dict(sorted(sources.items()))))
        SPONSORS_JSON.write_text(dump(sponsor_manifest(data['sponsors'])))
    return rep, data


def sponsor_manifest(sponsors):
    """Footer and poster logos (src/sponsors.json): sheet order, sizes from the files."""
    old = {s['id']: s for s in json.loads(SPONSORS_JSON.read_text())} if SPONSORS_JSON.exists() else {}
    out = []
    for s in sponsors:
        src = f'assets/sponsors/{s["id"]}.webp'
        with Image.open(ROOT / src) as im:
            w, h = im.size
        out.append({'id': s['id'], 'name': s['name'], 'src': src, 'w': w, 'h': h,
                    'scale': old.get(s['id'], {}).get('scale', 1.0)})
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument('--url', help='link to the Google Sheet (published or shared)')
    src.add_argument('--file', help='a local .xlsx copy of the sheet')
    ap.add_argument('--check', action='store_true', help='only check, write nothing')
    ap.add_argument('--expect-same', action='store_true', help='fail if the data differs from src/sheet.json')
    args = ap.parse_args()
    try:
        if args.url is not None:
            url = args.url.strip()
            if not url:
                raise SyncError('Tabulas saite nav iestatīta (GitHub: Settings → Secrets and variables → Actions → SHEET_URL).')
            try:
                body, _ = fetch(sheet_export_url(url), 'Tabula')
            except ValueError as e:
                raise SyncError(f'{e}. Pārbaudi, vai tabula joprojām ir publicēta '
                                '(Fails → Kopīgot → Publicēt tīmeklī) un saite ir pareiza.')
            if not body.startswith(b'PK'):
                raise SyncError('Pēc saites neatveras tabula. Pārbaudi, vai tā ir publicēta: '
                                'Fails → Kopīgot → Publicēt tīmeklī → „Viss dokuments”.')
            wb = load_workbook(io.BytesIO(body), data_only=True)
        else:
            wb = load_workbook(args.file, data_only=True)
        rep, data = sync(wb, write=not args.check and not args.expect_same)
        report = rep.markdown()
        code = 1 if rep.errors else 0
        if args.expect_same and data is not None and data != json.loads(SHEET_JSON.read_text()):
            report += '\nThe sheet does not match src/sheet.json: ' + '; '.join(rep.changes) + '\n'
            code = 3
    except (SyncError, ValueError) as e:
        report = f'## Tabulu neizdevās nolasīt — mājaslapa netika atjaunota\n\n{e}\n'
        code = 2
    print(report)
    summary = os.environ.get('GITHUB_STEP_SUMMARY')
    if summary:
        with open(summary, 'a', encoding='utf-8') as fh:
            fh.write(report)
    sys.exit(code)


if __name__ == '__main__':
    main()
