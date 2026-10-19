"""Builds the Google Sheets template (as .xlsx) for the 2026/27 season:
settings, teams, halls and sponsors for the site + the season plan.

    python3 tools/season_sheet.py <output.xlsx>

Then recalculate the formulas (LibreOffice) before sending the file.
"Iestatījumi", "Komandas", "Zāles" and "Sponsori" are prefilled from
src/sheet.json (what the site has now), so importing the template with
tools/sheet_sync.py gives back the same data.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Dobeles_basketbols_2026-27.xlsx'
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tools'))
from sheet_sync import SETTINGS, dig, site_logo_link  # noqa: E402  (same names as the importer)

SHEET = json.loads((ROOT / 'src' / 'sheet.json').read_text())
KG = 'Playoff kvalifikācijas grupa'

FONT = 'Arial'
NAVY = '112C5B'
INPUT = 'FFF6D5'      # light yellow: organisers fill these
AUTO = 'ECEFF3'       # light grey: prefilled / automatic, don't change
WARN = 'C0392B'
thin = Side(style='thin', color='D5D9E0')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(bold=False, color='1B2230', size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)

def header(ws, cols, widths):
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = f(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'

def style_rows(ws, first, last, ncols, kinds):
    """kinds: list per column — 'in' (to fill), 'auto' (prefilled/locked), 'calc' (formula)."""
    for r in range(first, last + 1):
        for ci in range(1, ncols + 1):
            c = ws.cell(row=r, column=ci)
            c.font = f()
            c.border = BORDER
            c.alignment = Alignment(vertical='center')
            k = kinds[ci - 1]
            if k == 'in':
                c.fill = PatternFill('solid', fgColor=INPUT)
            elif k in ('auto', 'calc'):
                c.fill = PatternFill('solid', fgColor=AUTO)
            if k == 'calc':
                c.font = f(color=WARN, bold=True)

def dv_list(ws, ref, rng, prompt=None, strict=True):
    dv = DataValidation(type='list', formula1=ref, allow_blank=True, showErrorMessage=strict)
    if strict:
        dv.errorTitle = 'Nederīga vērtība'
        dv.error = 'Izvēlies vērtību no saraksta.'
    if prompt:
        dv.promptTitle = 'Padoms'
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)

wb = Workbook()

# ---------------------------------------------------------------- Kā lietot
ws = wb.active
ws.title = 'Kā lietot'
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 100
ws.sheet_view.showGridLines = False
rows = [
    ('title', 'Dobeles basketbols 2026/27 — sezonas tabula'),
    ('text', 'No šīs tabulas mājaslapa ņem iestatījumus, komandas, zāles un sponsorus. Mājaslapa atjaunojas, kad GitHub palaiž „Atjaunot datus no tabulas” (sk. zemāk).'),
    ('gap', ''),
    ('h', 'Ko mājaslapa nolasa šobrīd'),
    ('kv', ('Jau tagad', 'Lapas „Iestatījumi”, „Komandas”, „Zāles” un „Sponsori”.')),
    ('kv', ('Vēl ne', 'Lapas „Kalendārs”, „MVP” un „Lēmumi”. Līdz tam mājaslapā redzamas demo spēles ar īstajām komandām un testa versijas paziņojums.')),
    ('kv', ('Vēlāk', 'Pēc izlozes — kalendārs no šīs tabulas. Kad būs LBS atļauja — rezultāti un statistika no basket.lv.')),
    ('gap', ''),
    ('h', 'Kā atjaunot mājaslapu'),
    ('kv', ('Poga', 'GitHub → dobelesbasketbols-site → Actions → „Atjaunot datus no tabulas” → Run workflow. Pēc 2–3 minūtēm izmaiņas ir mājaslapā.')),
    ('kv', ('Ja ir kļūdas', 'Mājaslapa paliek tāda, kāda bija, un GitHub atsūta e-pastu. Darbības lapā ir saraksts: lapa, rinda un kas jālabo.')),
    ('kv', ('Tikko laboji?', 'Google publicēto versiju atjauno dažu minūšu laikā. Ja izmaiņas neparādās, pagaidi un palaid vēlreiz.')),
    ('gap', ''),
    ('h', 'Logo (komandām un sponsoriem)'),
    ('kv', ('1. Mape', 'Google Drive izveido mapi (piem., „Logo”) un ielādē tur logo: PNG ar caurspīdīgu fonu. Der arī JPG vai WebP, SVG neder.')),
    ('kv', ('2. Piekļuve', 'Mapei: Kopīgot → Vispārējā piekļuve → „Ikviens, kam ir saite” → Skatītājs (Share → General access → Anyone with the link → Viewer).')),
    ('kv', ('3. Saite', 'Katram failam: Kopīgot → Kopēt saiti (Share → Copy link). Ielīmē to kolonnā „Logo (saite)” pie komandas vai sponsora.')),
    ('kv', ('Jau aizpildīts', 'Ja logo jau ir mājaslapā, kolonnā „Logo (saite)” ir tā adrese mājaslapā (…/assets/…webp). To neko nemaina — logo paliek.')),
    ('kv', ('Tukša šūna', 'Paliek logo, kas mājaslapā jau ir. Ja komandai logo nav, rāda iniciāļus.')),
    ('kv', ('Nomainīt logo', 'Adreses vietā ielīmē Google Drive saiti uz jauno failu un atjauno mājaslapu.')),
    ('gap', ''),
    ('h', 'Kā dati apvienosies sezonas laikā'),
    ('kv', ('basket.lv', 'Galvenais avots. Ja spēle ir basket.lv, mājaslapa ņem no turienes datumu, laiku, zāli, komandas, rezultātu un statistiku.')),
    ('kv', ('Šī tabula', 'Rezerves plāns visai sezonai (arī spēlēm, kuru komandas vēl nav zināmas) un manuālie dati: MVP, komiteju lēmumi, komandas, zāles.')),
    ('kv', ('Lēmumi', 'Tabulas lapa „Lēmumi” vienmēr ir svarīgāka par basket.lv (piem., neierašanās → 20:0).')),
    ('kv', ('Tabulas, izspēles shēma', 'Netiek rakstītas — mājaslapa tās aprēķina no rezultātiem.')),
    ('gap', ''),
    ('h', 'Krāsas'),
    ('fill_in', ('Dzeltens', 'Aizpilda organizatori.')),
    ('fill_auto', ('Pelēks', 'Sagatavots iepriekš vai aprēķināts automātiski — nemainīt.')),
    ('fill_calc', ('Sarkans teksts', 'Kolonna „Pārbaude” — ja tur ir teksts, rindā ir kļūda.')),
    ('gap', ''),
    ('h', 'Lapas'),
    ('kv', ('Iestatījumi', 'Sezona, kausu nosaukumi, MVP balvas sponsori, kājenes teksts un turnīra sociālo tīklu saites, sponsoru lapas teksts, basket.lv saites.')),
    ('kv', ('Komandas', 'Visas 12 komandas. Kolonnu „Grupa” aizpilda pēc izlozes. „basket.lv nosaukums” — tieši tā, kā komanda rakstīta basket.lv. Logo — saite uz failu Google Drive. Pilsēta, apraksts un saites nav obligāti — ja aizpildīti, tie redzami komandas lapā.')),
    ('kv', ('Zāles', 'Zāles ar adresēm. Spēlēs zāli norāda ar tās ID (dobele, jelgava).')),
    ('kv', ('Kalendārs', 'Visas sezonas spēles — viena rinda = viena spēle. Kārta = visas spēles vienā datumā un zālē (ne vairāk kā 4).')),
    ('kv', ('MVP', 'Katras kārtas MVP katram kausam. Statistika tiek ņemta no basket.lv automātiski.')),
    ('kv', ('Lēmumi', 'Neierašanās, anulēti rezultāti, labojumi, izlozes.')),
    ('kv', ('Sponsori', 'Sponsoru lapa un logo kājenē: līmenis (Ģenerālsponsors / Atbalstītājs), secība, apraksts, saites, logo. Tukša saite — poga netiek rādīta. Izdzēsta rinda — sponsors pazūd no mājaslapas.')),
    ('kv', ('Pārbaude', 'Kopsavilkums: cik spēļu ir ieplānotas un vai tabulā nav kļūdu.')),
    ('kv', ('Saraksti', 'Izvēlņu saraksti. Nemainīt.')),
    ('gap', ''),
    ('h', 'Vietu kodi (kad komandas vēl nav zināmas)'),
    ('kv', ('A1 … A6, B1 … B6', 'A vai B grupas 1.–6. vieta grupu turnīrā.')),
    ('kv', ('K1 … K8', 'Playoff kvalifikācijas grupas 1.–8. vieta (mājaslapā redzams kā PK1 … PK8).')),
    ('kv', ('W:SF1 / L:SF1', '1. pusfināla uzvarētājs / zaudētājs (W = uzvarētājs, L = zaudētājs). Tāpat W:H1, W:R16_1, W:QF2 utt.')),
    ('kv', ('Pāris', 'Izslēgšanas spēļu pāris: SF1, SF2 (pusfināli), BRONZE (bronzas fināls), GOLD (zelta fināls), H1, H2, HF (Cerību kauss), R16_1 … R16_4 (astotdaļfināli), QF1 … QF4 (ceturtdaļfināli), F (Izaicinājuma kausa fināls).')),
    ('kv', ('Kad komandas zināmas', 'Var atstāt kodu — mājaslapa pati ieliks komandu pēc tabulas. Vai arī ierakstīt komandas ID.')),
    ('gap', ''),
    ('h', 'Piemēru rindas (tikai paraugs — tabulā tās nav jāievada)'),
    ('ex', ('Komandas', 'buki · BK Buki / Sandramebeles.lv · Buki · BK Buki · A · https://drive.google.com/file/d/1AbC…/view · Dobele · Īss apraksts … · https://… · https://www.facebook.com/… · · ')),
    ('ex', ('Kalendārs', 'DBCK-GR-01 · DBČK · Grupu turnīrs · 1. kārta · — · — · 25.10.2026 · 11:00 · dobele · buki · doks · Plānota')),
    ('ex', ('Kalendārs (izslēgšana)', 'DBCK-SF1-3 · DBČK · Pusfināls · — · SF1 · 3 · 04.04.2027 · 14:00 · dobele · K1 · K4 · Ja vajag')),
    ('ex', ('MVP', 'DBČK · 25.10.2026 · Mārtiņš Ozols · skandijs')),
    ('ex', ('Lēmumi', 'DBCK-GR-07 · Neieradās viesi · — · 09.11.2026 · Sacensību komiteja · Komanda neieradās')),
    ('ex', ('Lēmumi (izloze)', 'DBCK-GR-30 · Izloze · buki, doks · 15.12.2026 · Sacensību komiteja · A grupas 3.–4. vieta')),
    ('ex', ('Sponsori', 'tervete · Alus darītava „Tērvete” · Atbalstītājs · 4 · Īss apraksts … · https://www.tervete.lv · https://www.facebook.com/… · · · https://…/assets/sponsors/tervete.webp')),
    ('gap', ''),
    ('h', 'Noteikumi, lai mājaslapa varētu nolasīt datus'),
    ('kv', ('1', 'Nemainīt lapu un kolonnu nosaukumus un to secību. Lapā „Iestatījumi” nemainīt kolonnu „Parametrs”.')),
    ('kv', ('ID', 'Komandu, zāļu un sponsoru ID: mazie latīņu burti bez garumzīmēm, cipari un domuzīme (piem., „spriditis”, „ozols-pub”). Pēc sezonas sākuma ID nemainīt.')),
    ('kv', ('Saites', 'Visas saites pilnas, sākas ar https://.')),
    ('kv', ('2', 'Datumu rakstīt kā DD.MM.GGGG, laiku kā HH:MM.')),
    ('kv', ('3', 'Spēles ID nemainīt — uz to atsaucas lapas „MVP” un „Lēmumi”. Jaunai spēlei pievieno jaunu rindu ar jaunu ID.')),
    ('kv', ('4', 'Spēli nedzēst — ja tā nenotiks, statusu maina uz „Atcelta”.')),
    ('kv', ('5', 'Vienā dienā viena zāle, ne vairāk kā 4 spēles (kolonna „Pārbaude” brīdina).')),
    ('gap', ''),
    ('h', 'Pēc augšupielādes Google Drive (vienreiz)'),
    ('kv', ('Atvērt', 'Augšupielādē failu Google Drive un atver ar Google Sheets (Fails → Saglabāt kā Google izklājlapu).')),
    ('kv', ('Aizsargāt', 'Dati → Aizsargāt lapas un diapazonus: aizsargā pelēkās kolonnas un lapu „Saraksti”, lai tās varētu mainīt tikai tu.')),
    ('kv', ('Publicēt', 'Fails → Kopīgot → Publicēt tīmeklī → „Viss dokuments” → Publicēt (File → Share → Publish to web → Entire document). Atstāj ieslēgtu automātisko atkārtoto publicēšanu. Saiti iedod Claude — mājaslapa no tās nolasīs datus.')),
    ('kv', ('Svarīgi', 'Publicētu tabulu var izlasīt ikviens, kam ir saite. Neraksti tajā tālruņus, e-pastus vai citus personas datus.')),
    ('kv', ('Kopīgot', 'Pārējiem organizatoriem dod rediģēšanas tiesības (Kopīgot → Redaktors).')),
]
r = 1
for kind, val in rows:
    if kind == 'title':
        c = ws.cell(row=r, column=1, value=val); c.font = f(bold=True, size=16, color=NAVY)
        ws.row_dimensions[r].height = 26
    elif kind == 'text':
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=val); c.font = f(); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
    elif kind == 'h':
        c = ws.cell(row=r, column=1, value=val); c.font = f(bold=True, size=12, color=NAVY)
        ws.row_dimensions[r].height = 20
    elif kind in ('kv', 'ex', 'fill_in', 'fill_auto', 'fill_calc'):
        k, v = val
        a = ws.cell(row=r, column=1, value=k); b = ws.cell(row=r, column=2, value=v)
        a.font = f(bold=True); b.font = f(italic=(kind == 'ex'))
        a.alignment = Alignment(vertical='top', wrap_text=True); b.alignment = Alignment(vertical='top', wrap_text=True)
        if kind == 'fill_in':
            a.fill = PatternFill('solid', fgColor=INPUT)
        elif kind == 'fill_auto':
            a.fill = PatternFill('solid', fgColor=AUTO)
        elif kind == 'fill_calc':
            a.fill = PatternFill('solid', fgColor=AUTO); a.font = f(bold=True, color=WARN)
        lines = max(1, len(v) // 95 + 1)
        ws.row_dimensions[r].height = 15 * lines + 2
    r += 1

# ---------------------------------------------------------------- Iestatījumi
ws = wb.create_sheet('Iestatījumi')
header(ws, ['Parametrs', 'Vērtība', 'Paskaidrojums'], [30, 60, 62])
HELP = {
    'Sezona': 'Sezonas nosaukums mājaslapā, piem., 2026/27.',
    'DBČK nosaukums': 'Pilns kausa nosaukums (sākumlapas virsraksts, kopvērtējums).',
    'JSSC nosaukums': 'Pilns kausa nosaukums (sākumlapas virsraksts, kopvērtējums).',
    'DBČK MVP balvas sponsors': 'Nosaukums tieši kā lapā „Sponsori”. Ja maini sponsoru, atsūti Claude baltu logo versiju MVP blokam.',
    'JSSC MVP balvas sponsors': 'Nosaukums tieši kā lapā „Sponsori”. Ja maini sponsoru, atsūti Claude baltu logo versiju MVP blokam.',
    'DBČK basket.lv saite': 'Saite uz Dobeles kausa sacensībām basket.lv. Vēl netiek izmantota.',
    'JSSC basket.lv saite': 'Saite uz Izaicinājuma kausa sacensībām basket.lv. Vēl netiek izmantota.',
    'Sponsoru lapas teksts': 'Pateicības teksts lapas „Sponsori” augšā. Tukšs — teksta nav.',
    'Kājenes teksts': 'Teksts lapas apakšā (kas organizē sacensības). Tukšs — teksta nav.',
    'Turnīra Facebook': 'Turnīra lapa (pilna saite ar https://). Tukša — ikona kājenē netiek rādīta.',
    'Turnīra Instagram': 'Turnīra konts (pilna saite ar https://). Tukša — ikona netiek rādīta.',
    'Turnīra YouTube': 'Turnīra kanāls (pilna saite ar https://). Tukša — ikona netiek rādīta.',
    'Turnīra TikTok': 'Turnīra konts (pilna saite ar https://). Tukša — ikona netiek rādīta.',
}
settings = [(param, dig(SHEET['settings'], key) or None, HELP[param]) for param, key, *_ in SETTINGS]
for i, row in enumerate(settings, start=2):
    for j, v in enumerate(row, start=1):
        ws.cell(row=i, column=j, value=v)
style_rows(ws, 2, 1 + len(settings), 3, ['auto', 'in', 'auto'])
for i in range(2, 2 + len(settings)):
    ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=3).font = f(color='5B6478')
    v = ws.cell(row=i, column=2).value or ''
    ws.row_dimensions[i].height = max(30, 15 * (len(str(v)) // 60 + 1) + 4)
    ws.cell(row=i, column=2).number_format = '@'   # "2026/27" must stay text
ws.cell(row=1, column=1).comment = Comment('Parametru nosaukumus nemainīt — mājaslapa tos meklē pēc nosaukuma.', 'Claude')

# ---------------------------------------------------------------- Komandas
ws = wb.create_sheet('Komandas')
# Columns G–L are optional (shown on the team page only when filled).
header(ws, ['ID', 'Pilnais nosaukums', 'Īsais nosaukums', 'basket.lv nosaukums', 'Grupa', 'Logo (saite)',
            'Pilsēta', 'Apraksts', 'Mājaslapa', 'Facebook', 'Instagram', 'TikTok', 'Piezīme', 'Pārbaude'],
       [14, 34, 18, 30, 9, 34, 16, 50, 28, 28, 28, 28, 30, 44])
TEAM_ROWS = 20   # room for changes
for i, t in enumerate(SHEET['teams'], start=2):
    has_logo = (ROOT / 'assets' / 'teams' / f'{t["id"]}.webp').exists()
    vals = [t['id'], t['name'], t['short'], '; '.join(t.get('basket', [])) or None, t.get('group') or None,
            site_logo_link('teams', t['id']) if has_logo else None, t.get('city'), t.get('description'), t.get('website'), t.get('facebook'), t.get('instagram'), t.get('tiktok')]
    for j, v in enumerate(vals, start=1):
        ws.cell(row=i, column=j, value=v)
    if not has_logo:
        ws.cell(row=i, column=13, value='Logo vēl nav — ieliec saiti kolonnā „Logo (saite)”.')
style_rows(ws, 2, 1 + TEAM_ROWS, 14, ['in'] * 13 + ['calc'])
for r in range(2, 2 + TEAM_ROWS):
    R_ = str(r)
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical='top')
    for ci in (1, 3, 4, 5):
        ws.cell(row=r, column=ci).number_format = '@'
    link_checks = ''.join(
        f'IF(AND({c}{R_}<>"",LEFT({c}{R_},8)<>"https://"),"{name}: saitei jāsākas ar https://. ","")&'
        for c, name in (('F', 'Logo'), ('I', 'Mājaslapa'), ('J', 'Facebook'), ('K', 'Instagram'), ('L', 'TikTok')))
    ws.cell(row=r, column=14, value=(
        f'=IF(AND(A{R_}="",B{R_}="",C{R_}=""),"",TRIM('
        f'IF(A{R_}="","Nav ID. ","")&'
        f'IF(AND(A{R_}<>"",OR(NOT(EXACT(A{R_},LOWER(A{R_}))),ISNUMBER(FIND(" ",A{R_})))),"ID: tikai mazie burti, bez atstarpēm. ","")&'
        f'IF(B{R_}="","Nav pilnā nosaukuma. ","")&'
        f'IF(C{R_}="","Nav īsā nosaukuma. ",IF(LEN(C{R_})>12,"Īsais nosaukums garāks par 12 zīmēm. ",""))&'
        f'IF(AND(E{R_}<>"",E{R_}<>"A",E{R_}<>"B"),"Grupa: A vai B. ","")&'
        + link_checks +
        f'IF(AND(A{R_}<>"",COUNTIF($A$2:$A${1 + TEAM_ROWS},A{R_})>1),"Atkārtots ID. ","")'
        f'))'))
ws.cell(row=1, column=1).comment = Comment('Īss ID: mazie latīņu burti bez garumzīmēm, cipari, domuzīme. Pēc sezonas sākuma nemainīt — uz to atsaucas kalendārs.', 'Claude')
ws.cell(row=1, column=3).comment = Comment('Rāda tabulās, kalendārā un plakātos. Līdz 12 zīmēm.', 'Claude')
ws.cell(row=1, column=4).comment = Comment('Tieši tā, kā komanda rakstīta basket.lv. Ja nosaukumi ir vairāki, atdali ar semikolu (;).', 'Claude')
ws.cell(row=1, column=5).comment = Comment('A vai B — aizpilda pēc izlozes. Kamēr nav 6 + 6, demo spēlēs grupas ir simulētas.', 'Claude')
ws.cell(row=1, column=6).comment = Comment('Saite uz logo. Ja logo jau ir mājaslapā, šeit ir tā adrese mājaslapā — neko nemainīt. Jaunam vai citam logo: Google Drive faila saite (Kopīgot → Kopēt saiti; piekļuve „Ikviens, kam ir saite”), PNG ar caurspīdīgu fonu.', 'Claude')
ws.cell(row=1, column=7).comment = Comment('Nav obligāti. Pilsēta vai novads komandas lapā.', 'Claude')
ws.cell(row=1, column=8).comment = Comment('Nav obligāti. Īss apraksts komandas lapā (2–4 teikumi).', 'Claude')
ws.cell(row=1, column=9).comment = Comment('Nav obligāti. Pilna saite, sākas ar https://. Tukša šūna — poga netiek rādīta.', 'Claude')
ws.cell(row=1, column=14).comment = Comment('Aprēķinās automātiski. Ja šeit parādās teksts, rindā ir kļūda.', 'Claude')
dv_list(ws, '"A,B"', f'E2:E{1 + TEAM_ROWS}')

# ---------------------------------------------------------------- Zāles
ws = wb.create_sheet('Zāles')
header(ws, ['ID', 'Nosaukums', 'Adrese', 'Piezīme'], [14, 30, 34, 40])
HALL_ROWS = 6
for i, h in enumerate(SHEET['venues'], start=2):
    note = 'Nosaukums jāprecizē.' if h['name'] == 'Jelgavas Sporta halle' else None
    for j, v in enumerate((h['id'], h['name'], h.get('address'), note), start=1):
        ws.cell(row=i, column=j, value=v)
style_rows(ws, 2, 1 + HALL_ROWS, 4, ['in', 'in', 'in', 'in'])
ws.cell(row=1, column=1).comment = Comment('Īss ID (mazie burti). Kalendārā zāli norāda ar šo ID.', 'Claude')

# ---------------------------------------------------------------- Kalendārs
ws = wb.create_sheet('Kalendārs')
cols = ['Spēles ID', 'Kauss', 'Posms', 'Kārta', 'Pāris', 'Spēle Nr.', 'Datums', 'Laiks', 'Zāle',
        'Mājinieki', 'Viesi', 'Statuss', 'basket.lv saite', 'Piezīme', 'Pārbaude']
header(ws, cols, [15, 8, 28, 11, 9, 9, 12, 8, 11, 13, 13, 12, 30, 28, 46])
games = []   # (id, cup, stage, round, tie, n, home, away, status)
for i in range(1, 31):
    games.append((f'DBCK-GR-{i:02d}', 'DBČK', 'Grupu turnīrs', '', '', '', '', '', 'Plānota'))
k = 1
for a in range(1, 5):
    for b in range(1, 5):
        games.append((f'DBCK-K-{k:02d}', 'DBČK', KG, '', '', '', f'A{a}', f'B{b}', 'Plānota')); k += 1
games += [
    ('DBCK-H1', 'DBČK', 'Cerību kauss', '', 'H1', 1, 'A5', 'B6', 'Plānota'),
    ('DBCK-H2', 'DBČK', 'Cerību kauss', '', 'H2', 1, 'B5', 'A6', 'Plānota'),
    ('DBCK-HF', 'DBČK', 'Cerību kauss', '', 'HF', 1, 'W:H1', 'W:H2', 'Plānota'),
]
for tie, hi, lo in (('SF1', 'K1', 'K4'), ('SF2', 'K2', 'K3')):
    games += [(f'DBCK-{tie}-1', 'DBČK', 'Pusfināls', '', tie, 1, hi, lo, 'Plānota'),
              (f'DBCK-{tie}-2', 'DBČK', 'Pusfināls', '', tie, 2, lo, hi, 'Plānota'),
              (f'DBCK-{tie}-3', 'DBČK', 'Pusfināls', '', tie, 3, hi, lo, 'Ja vajag')]
games += [('DBCK-BRONZE', 'DBČK', 'Spēle par 3. vietu', '', 'BRONZE', 1, 'L:SF1', 'L:SF2', 'Plānota'),
          ('DBCK-GOLD', 'DBČK', 'Fināls', '', 'GOLD', 1, 'W:SF1', 'W:SF2', 'Plānota')]
R16 = [('R16_1', 'B4', 'A5'), ('R16_2', 'A3', 'B6'), ('R16_3', 'A4', 'B5'), ('R16_4', 'B3', 'A6')]
for tie, a, b in R16:
    games += [(f'JSSC-{tie}-1', 'JSSC', 'Astotdaļfināls', '', tie, 1, a, b, 'Plānota'),
              (f'JSSC-{tie}-2', 'JSSC', 'Astotdaļfināls', '', tie, 2, b, a, 'Plānota')]
QF = [('QF1', 'A1', 'W:R16_1'), ('QF2', 'B2', 'W:R16_2'), ('QF3', 'B1', 'W:R16_3'), ('QF4', 'A2', 'W:R16_4')]
for tie, seed, w in QF:
    games += [(f'JSSC-{tie}-1', 'JSSC', 'Ceturtdaļfināls', '', tie, 1, w, seed, 'Plānota'),
              (f'JSSC-{tie}-2', 'JSSC', 'Ceturtdaļfināls', '', tie, 2, seed, w, 'Plānota')]
# Challenge Cup semifinals and final: one game each.
for tie, a, b in (('SF1', 'W:QF1', 'W:QF2'), ('SF2', 'W:QF3', 'W:QF4')):
    games.append((f'JSSC-{tie}', 'JSSC', 'Pusfināls', '', tie, 1, a, b, 'Plānota'))
games.append(('JSSC-F', 'JSSC', 'Fināls', '', 'F', 1, 'W:SF1', 'W:SF2', 'Plānota'))

CAL_ROWS = 160   # prefilled games + room for extra rows
for i, (gid, cup, stage, rnd, tie, n, home, away, status) in enumerate(games, start=2):
    vals = [gid, cup, stage, rnd or None, tie or None, n or None, None, None, None, home or None, away or None, status]
    for j, v in enumerate(vals, start=1):
        ws.cell(row=i, column=j, value=v)
last = 1 + CAL_ROWS
kinds = ['auto', 'auto', 'auto', 'in', 'auto', 'auto', 'in', 'in', 'in', 'in', 'in', 'in', 'in', 'in', 'calc']
style_rows(ws, 2, last, 15, kinds)
# Extra rows below the prefilled season are fully editable.
for r in range(2 + len(games), last + 1):
    for ci in range(1, 7):
        ws.cell(row=r, column=ci).fill = PatternFill('solid', fgColor=INPUT)
for r in range(2, last + 1):
    ws.cell(row=r, column=7).number_format = 'DD.MM.YYYY'
    ws.cell(row=r, column=8).number_format = 'HH:MM'
    ws.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='center')
    R = str(r)
    G, H, I, J, K, L, A = 'G' + R, 'H' + R, 'I' + R, 'J' + R, 'K' + R, 'L' + R, 'A' + R
    rng = lambda col: f'${col}$2:${col}${last}'
    ws.cell(row=r, column=15, value=(
        f'=IF({A}="","",TRIM('
        f'IF(AND({J}<>"",ISNA(MATCH({J},Saraksti!$A$2:$A$90,0))),"Nezināma komanda: "&{J}&". ","")&'
        f'IF(AND({K}<>"",ISNA(MATCH({K},Saraksti!$A$2:$A$90,0))),"Nezināma komanda: "&{K}&". ","")&'
        f'IF(AND({J}<>"",{J}={K}),"Mājinieki un viesi ir viena komanda. ","")&'
        f'IF(AND({G}<>"",{I}=""),"Nav norādīta zāle. ","")&'
        f'IF(AND({G}<>"",{I}<>"",COUNTIFS({rng("G")},{G},{rng("I")},{I},{rng("L")},"<>Atcelta")>4),"Vairāk nekā 4 spēles šajā dienā. ","")&'
        f'IF(AND({G}<>"",{I}<>"",COUNTIFS({rng("G")},{G},{rng("I")},"<>"&{I},{rng("I")},"<>",{rng("L")},"<>Atcelta")>0),"Šajā dienā spēles vairākās zālēs. ","")&'
        f'IF(AND({G}<>"",{H}<>"",COUNTIFS({rng("G")},{G},{rng("I")},{I},{rng("H")},{H},{rng("L")},"<>Atcelta")>1),"Divas spēles vienā laikā. ","")&'
        f'IF(AND({G}<>"",{J}<>"",{L}<>"Atcelta",COUNTIFS({rng("G")},{G},{rng("J")},{J},{rng("L")},"<>Atcelta")+COUNTIFS({rng("G")},{G},{rng("K")},{J},{rng("L")},"<>Atcelta")>1),{J}&" spēlē divreiz vienā dienā. ","")&'
        f'IF(AND({G}<>"",{K}<>"",{K}<>{J},{L}<>"Atcelta",COUNTIFS({rng("G")},{G},{rng("J")},{K},{rng("L")},"<>Atcelta")+COUNTIFS({rng("G")},{G},{rng("K")},{K},{rng("L")},"<>Atcelta")>1),{K}&" spēlē divreiz vienā dienā. ","")&'
        f'IF(COUNTIF({rng("A")},{A})>1,"Atkārtots spēles ID. ","")'
        f'))'))
ws.auto_filter.ref = f'A1:O{last}'
dv_list(ws, '=Saraksti!$B$2:$B$3', f'B2:B{last}')
dv_list(ws, '=Saraksti!$C$2:$C$10', f'C2:C{last}')
dv_list(ws, '=Saraksti!$F$2:$F$9', f'D2:D{last}', prompt='Grupu turnīram un Playoff kvalifikācijas grupai: kārtas numurs.')
dv_list(ws, '=Saraksti!$E$2:$E$20', f'E2:E{last}')
dv_list(ws, '=Saraksti!$H$2:$H$7', f'I2:I{last}')
dv_list(ws, '=Saraksti!$A$2:$A$90', f'J2:K{last}', prompt='Komandas ID (piem., buki) vai vietas kods (A1, K4, W:SF1).', strict=False)
dv_list(ws, '=Saraksti!$D$2:$D$5', f'L2:L{last}')
dv_date = DataValidation(type='date', operator='between', formula1='DATE(2026,9,1)', formula2='DATE(2027,6,30)',
                         allow_blank=True, showErrorMessage=True, errorTitle='Datums', error='Datums jāraksta DD.MM.GGGG (sezonas laikā).')
ws.add_data_validation(dv_date); dv_date.add(f'G2:G{last}')
ws.cell(row=1, column=15).comment = Comment('Aprēķinās automātiski. Ja šeit parādās teksts, rindā ir kļūda.', 'Claude')
ws.cell(row=1, column=10).comment = Comment('Komandas ID no lapas „Komandas” vai vietas kods (A1, K4, W:SF1, L:SF2 …).', 'Claude')
ws.cell(row=1, column=13).comment = Comment('Nav obligāti. Aizpilda tikai, ja mājaslapa sasaista spēli ar nepareizu basket.lv spēli.', 'Claude')

# ---------------------------------------------------------------- MVP
ws = wb.create_sheet('MVP')
header(ws, ['Kauss', 'Kārtas datums', 'Spēlētājs', 'Komanda', 'Spēles ID', 'Piezīme'], [9, 14, 28, 14, 16, 36])
MVP_ROWS = 60
style_rows(ws, 2, 1 + MVP_ROWS, 6, ['in'] * 6)
for r in range(2, 2 + MVP_ROWS):
    ws.cell(row=r, column=2).number_format = 'DD.MM.YYYY'
dv_list(ws, '=Saraksti!$B$2:$B$3', f'A2:A{1 + MVP_ROWS}')
dv_list(ws, '=Saraksti!$A$2:$A$21', f'D2:D{1 + MVP_ROWS}')
dv_list(ws, f'=Kalendārs!$A$2:$A${last}', f'E2:E{1 + MVP_ROWS}', prompt='Nav obligāti: spēle, kurā MVP spēlēja.', strict=False)
ws.cell(row=1, column=3).comment = Comment('Vārds un uzvārds tieši tā, kā basket.lv protokolā — tad statistika pievienosies automātiski.', 'Claude')

# ---------------------------------------------------------------- Lēmumi
ws = wb.create_sheet('Lēmumi')
header(ws, ['Spēles ID', 'Veids', 'Vērtība', 'Lēmuma datums', 'Kas pieņēma', 'Piezīme'], [16, 22, 24, 14, 22, 44])
DEC_ROWS = 40
style_rows(ws, 2, 1 + DEC_ROWS, 6, ['in'] * 6)
for r in range(2, 2 + DEC_ROWS):
    ws.cell(row=r, column=4).number_format = 'DD.MM.YYYY'
dv_list(ws, f'=Kalendārs!$A$2:$A${last}', f'A2:A{1 + DEC_ROWS}')
dv_list(ws, '=Saraksti!$G$2:$G$6', f'B2:B{1 + DEC_ROWS}')
ws.cell(row=1, column=3).comment = Comment('Labojumam: rezultāts (78:71). Izlozei: komandu ID secībā (buki, doks). Citiem — tukšs.', 'Claude')

# ---------------------------------------------------------------- Sponsori
# One row per sponsor: the site's "Sponsori" page and the footer logos.
# Existing sponsors already have a logo on the site (matched by ID); the logo
# link is only needed for a new or changed logo. An empty link cell = no button.
ws = wb.create_sheet('Sponsori')
header(ws, ['ID', 'Nosaukums', 'Līmenis', 'Secība', 'Apraksts', 'Mājaslapa', 'Facebook', 'Instagram', 'TikTok', 'Logo (saite)', 'Pārbaude'],
       [20, 30, 17, 8, 60, 30, 30, 30, 30, 34, 44])
SP_ROWS = 30
for i, row in enumerate(SHEET['sponsors'], start=2):
    vals = [row['id'], row['name'], row['level'], row['order'], row['description'] or None,
            row['website'] or None, row['facebook'] or None, row['instagram'] or None, row['tiktok'] or None,
            site_logo_link('sponsors', row['id']) if (ROOT / 'assets' / 'sponsors' / f'{row["id"]}.webp').exists() else None]
    for j, v in enumerate(vals, start=1):
        ws.cell(row=i, column=j, value=v)
style_rows(ws, 2, 1 + SP_ROWS, 11, ['auto', 'in', 'in', 'in', 'in', 'in', 'in', 'in', 'in', 'in', 'calc'])
for r in range(2, 2 + SP_ROWS):
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 30
    ws.cell(row=r, column=4).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=r, column=1).number_format = '@'
    R_ = str(r)
    link_checks = ''.join(
        f'IF(AND({c}{R_}<>"",LEFT({c}{R_},8)<>"https://"),"{name}: saitei jāsākas ar https://. ","")&'
        for c, name in (('F', 'Mājaslapa'), ('G', 'Facebook'), ('H', 'Instagram'), ('I', 'TikTok'), ('J', 'Logo')))
    ws.cell(row=r, column=11, value=(
        f'=IF(AND(A{R_}="",B{R_}=""),"",TRIM('
        f'IF(A{R_}="","Nav ID. ","")&'
        f'IF(AND(A{R_}<>"",OR(NOT(EXACT(A{R_},LOWER(A{R_}))),ISNUMBER(FIND(" ",A{R_})))),"ID: tikai mazie burti, bez atstarpēm. ","")&'
        f'IF(B{R_}="","Nav nosaukuma. ","")&'
        f'IF(ISNA(MATCH(C{R_},Saraksti!$I$2:$I$3,0)),"Izvēlies līmeni. ","")&'
        f'IF(AND(D{R_}<>"",NOT(ISNUMBER(D{R_}))),"Secībai jābūt skaitlim. ","")&'
        + link_checks +
        f'IF(AND(A{R_}<>"",COUNTIF($A$2:$A${1 + SP_ROWS},A{R_})>1),"Atkārtots ID. ","")'
        f'))'))
# New sponsor rows below the prefilled ones: ID is editable there.
for r in range(2 + len(SHEET['sponsors']), 2 + SP_ROWS):
    ws.cell(row=r, column=1).fill = PatternFill('solid', fgColor=INPUT)
dv_list(ws, '=Saraksti!$I$2:$I$3', f'C2:C{1 + SP_ROWS}')
ws.cell(row=1, column=1).comment = Comment('Saista rindu ar logo mājaslapā — esošajiem nemainīt. Jaunam sponsoram: ID (mazie burti, piem., „ozols-pub”) un logo saite.', 'Claude')
ws.cell(row=1, column=4).comment = Comment('Secība lapā līmeņa ietvaros: 1 = pirmais. Tāda pati secība kājenē.', 'Claude')
ws.cell(row=1, column=5).comment = Comment('Īss apraksts uznirstošajā logā (2–4 teikumi).', 'Claude')
ws.cell(row=1, column=6).comment = Comment('Pilna saite, sākas ar https://. Tukša šūna — poga netiek rādīta.', 'Claude')
ws.cell(row=1, column=10).comment = Comment('Saite uz logo. Ja logo jau ir mājaslapā, šeit ir tā adrese mājaslapā — neko nemainīt. Jaunam vai citam logo: Google Drive faila saite (Kopīgot → Kopēt saiti; piekļuve „Ikviens, kam ir saite”).', 'Claude')
ws.cell(row=1, column=11).comment = Comment('Aprēķinās automātiski. Ja šeit parādās teksts, rindā ir kļūda.', 'Claude')

# ---------------------------------------------------------------- Pārbaude
ws = wb.create_sheet('Pārbaude')
ws.sheet_view.showGridLines = False
for col, w in zip('ABCDE', (34, 16, 20, 16, 40)):
    ws.column_dimensions[col].width = w
ws['A1'] = 'Pārbaude'; ws['A1'].font = f(bold=True, size=16, color=NAVY)
ws['A3'] = 'Rindas ar kļūdām kalendārā'; ws['A3'].font = f(bold=True)
ws['B3'] = f'=COUNTIF(Kalendārs!$O$2:$O${last},"?*")'
ws['B3'].font = f(bold=True, color=WARN, size=12)
ws['C3'] = 'Ja nav 0 — skaties kolonnu „Pārbaude” lapā „Kalendārs”.'; ws['C3'].font = f(color='5B6478')
ws['A4'] = 'Komandas bez grupas'; ws['A4'].font = f(bold=True)
ws['B4'] = '=COUNTIFS(Komandas!$A$2:$A$21,"?*",Komandas!$E$2:$E$21,"")'
ws['A5'] = 'A grupā / B grupā'; ws['A5'].font = f(bold=True)
ws['B5'] = '=COUNTIF(Komandas!$E$2:$E$21,"A")&" / "&COUNTIF(Komandas!$E$2:$E$21,"B")'
ws['C5'] = 'Jābūt 6 / 6.'; ws['C5'].font = f(color='5B6478')
for c in ('B4', 'B5'):
    ws[c].font = f(bold=True, size=12)
hdr = ['Posms', 'Kauss', 'Spēles pēc nolikuma', 'Kalendārā', 'Ar datumu']
for j, h in enumerate(hdr, start=1):
    c = ws.cell(row=7, column=j, value=h)
    c.font = f(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY); c.border = BORDER
plan = [('Grupu turnīrs', 'DBČK', 30), (KG, 'DBČK', 16), ('Cerību kauss', 'DBČK', 3),
        ('Pusfināls', 'DBČK', '4–6'), ('Spēle par 3. vietu', 'DBČK', 1), ('Fināls', 'DBČK', 1),
        ('Astotdaļfināls', 'JSSC', 8), ('Ceturtdaļfināls', 'JSSC', 8), ('Pusfināls', 'JSSC', 2), ('Fināls', 'JSSC', 1)]
for i, (stage, cup, n) in enumerate(plan, start=8):
    ws.cell(row=i, column=1, value=stage); ws.cell(row=i, column=2, value=cup); ws.cell(row=i, column=3, value=n)
    ws.cell(row=i, column=4, value=f'=COUNTIFS(Kalendārs!$C$2:$C${last},A{i},Kalendārs!$B$2:$B${last},B{i},Kalendārs!$L$2:$L${last},"<>Atcelta")')
    ws.cell(row=i, column=5, value=f'=COUNTIFS(Kalendārs!$C$2:$C${last},A{i},Kalendārs!$B$2:$B${last},B{i},Kalendārs!$L$2:$L${last},"<>Atcelta",Kalendārs!$G$2:$G${last},"<>")')
    for j in range(1, 6):
        c = ws.cell(row=i, column=j); c.font = f(); c.border = BORDER
        c.alignment = Alignment(horizontal='left' if j <= 2 else 'center')
note_row = 8 + len(plan) + 1
ws.cell(row=note_row, column=1, value='DBČK pusfinālos 3. spēle notiek tikai, ja vajag (sērija līdz divām uzvarām). JSSC pusfināli un fināls — viena spēle.').font = f(color='5B6478', italic=True)
ws.cell(row=note_row + 2, column=1, value='Sponsoru rindas ar kļūdām').font = f(bold=True)
ws.cell(row=note_row + 2, column=2, value='=COUNTIF(Sponsori!$K$2:$K$31,"?*")').font = f(bold=True, color=WARN, size=12)
ws.cell(row=note_row + 2, column=3, value='Ja nav 0 — skaties kolonnu „Pārbaude” lapā „Sponsori”.').font = f(color='5B6478')
ws.cell(row=note_row + 3, column=1, value='Komandu rindas ar kļūdām').font = f(bold=True)
ws.cell(row=note_row + 3, column=2, value='=COUNTIF(Komandas!$N$2:$N$21,"?*")').font = f(bold=True, color=WARN, size=12)
ws.cell(row=note_row + 3, column=3, value='Ja nav 0 — skaties kolonnu „Pārbaude” lapā „Komandas”.').font = f(color='5B6478')

# ---------------------------------------------------------------- Saraksti
ws = wb.create_sheet('Saraksti')
header(ws, ['Komandas un vietu kodi', 'Kauss', 'Posms', 'Statuss', 'Pāris', 'Kārta', 'Lēmuma veids', 'Zāles', 'Sponsora līmenis'],
       [24, 10, 28, 12, 10, 12, 22, 12, 18])
for r in range(2, 22):   # team IDs come from the Komandas tab
    ws.cell(row=r, column=1, value=f'=IF(Komandas!A{r}="","",Komandas!A{r})')
codes = [f'A{i}' for i in range(1, 7)] + [f'B{i}' for i in range(1, 7)] + [f'K{i}' for i in range(1, 9)] + \
        ['W:H1', 'W:H2', 'W:SF1', 'L:SF1', 'W:SF2', 'L:SF2'] + [f'W:R16_{i}' for i in range(1, 5)] + [f'W:QF{i}' for i in range(1, 5)]
for i, code in enumerate(codes, start=22):
    ws.cell(row=i, column=1, value=code)
lists = {
    2: ['DBČK', 'JSSC'],
    3: ['Grupu turnīrs', KG, 'Cerību kauss', 'Pusfināls', 'Spēle par 3. vietu', 'Fināls', 'Astotdaļfināls', 'Ceturtdaļfināls'],
    4: ['Plānota', 'Ja vajag', 'Pārcelta', 'Atcelta'],
    5: ['SF1', 'SF2', 'BRONZE', 'GOLD', 'H1', 'H2', 'HF', 'R16_1', 'R16_2', 'R16_3', 'R16_4', 'QF1', 'QF2', 'QF3', 'QF4', 'F'],
    6: [f'{i}. kārta' for i in range(1, 9)],
    7: ['Neieradās mājinieki', 'Neieradās viesi', 'Rezultāts anulēts', 'Rezultāta labojums', 'Izloze'],
    9: ['Ģenerālsponsors', 'Atbalstītājs'],
}
for col, vals in lists.items():
    for i, v in enumerate(vals, start=2):
        ws.cell(row=i, column=col, value=v)
for r in range(2, 8):
    ws.cell(row=r, column=8, value=f'=IF(Zāles!A{r}="","",Zāles!A{r})')
for r in range(2, 2 + max(len(codes) + 20, 20)):
    for ci in range(1, 10):
        c = ws.cell(row=r, column=ci); c.font = f()
        if c.value is not None:
            c.fill = PatternFill('solid', fgColor=AUTO); c.border = BORDER

# Tab colours: inputs yellow-ish, reference grey
for name, color in (('Kā lietot', NAVY), ('Kalendārs', 'E9B716'), ('MVP', 'E9B716'), ('Lēmumi', 'E9B716'),
                    ('Komandas', 'E9B716'), ('Zāles', 'E9B716'), ('Iestatījumi', 'E9B716'), ('Sponsori', 'E9B716'),
                    ('Pārbaude', 'C0392B'), ('Saraksti', '8A909C')):
    wb[name].sheet_properties.tabColor = color

wb.save(OUT)
print('saved', OUT, len(games), 'games')
